"""
Build the IEEE Conference Paper for SmartDetect (Exact 6-Page Conference Target).

Peer-Review Ready IEEE Conference Format:
- Author Block: Exact 4 authors, affiliations, and emails.
- No placeholder conference headers.
- Strict IEEE structure, concise technical prose, no marketing or exaggerated novelty.
- Exact ChokePoint P1E_S1 experimental results preserved.
"""
from __future__ import annotations
import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BASE_DIR = Path(__file__).resolve().parent
FIGS = BASE_DIR / "figures"
OUT = BASE_DIR / "SmartDetect_IEEE_Paper.docx"

INK = RGBColor(0x14, 0x16, 0x1A)
MUT = RGBColor(0x50, 0x55, 0x60)
ACC = RGBColor(0x1F, 0x3A, 0x68)
OXI = RGBColor(0x8C, 0x2F, 0x21)

BODY_PT = 9.5
SERIF = "Times New Roman"
SANS = "Arial"

doc = Document()

# ── Page Setup: Standard IEEE margins (Top: 0.75", Bottom: 1.0", Left/Right: 0.62") ──
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11.0)
for attr, val in (("top_margin", 0.75), ("bottom_margin", 1.0),
                  ("left_margin", 0.62), ("right_margin", 0.62)):
    setattr(s, attr, Inches(val))

st = doc.styles["Normal"]
st.font.name = SERIF
st.font.size = Pt(BODY_PT)
st.paragraph_format.space_after = Pt(0)
st.paragraph_format.line_spacing = 1.0
st._element.rPr.rFonts.set(qn("w:eastAsia"), SERIF)


def set_cols(section, n, space_tw=280):
    sectPr = section._sectPr
    t = sectPr.find(qn("w:type"))
    if t is None:
        t = OxmlElement("w:type")
        sectPr.insert(0, t)
    t.set(qn("w:val"), "continuous")
    cols = sectPr.xpath("./w:cols")
    c = cols[0] if cols else OxmlElement("w:cols")
    if not cols:
        sectPr.append(c)
    c.set(qn("w:num"), str(n))
    c.set(qn("w:space"), str(space_tw))
    c.set(qn("w:equalWidth"), "1")


def new_section(n_cols):
    sec = doc.add_section(WD_SECTION.CONTINUOUS)
    sec.page_width, sec.page_height = Inches(8.5), Inches(11.0)
    sec.top_margin, sec.bottom_margin = Inches(0.75), Inches(1.0)
    sec.left_margin, sec.right_margin = Inches(0.62), Inches(0.62)
    set_cols(sec, n_cols)
    return sec


def para(text="", *, size=BODY_PT, bold=False, italic=False, align=None,
         font=SERIF, color=INK, before=0, after=0, indent=None, spacing=1.0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = spacing
    if indent is not None:
        pf.first_line_indent = Inches(indent)
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.font.name, r.font.size = font, Pt(size)
        r.bold, r.italic = bold, italic
        r.font.color.rgb = color
    return p


def rich(p, parts, size=BODY_PT, font=SERIF, color=INK):
    for text, sty in parts:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = "b" in sty
        r.italic = "i" in sty
        if "m" in sty:
            r.font.name = "Consolas"
            r.font.size = Pt(size - 0.8)
            r.font.color.rgb = ACC
        else:
            r.font.name = font
            r.font.color.rgb = OXI if "x" in sty else (ACC if "a" in sty else color)
    return p


def body(parts, indent=0.14, after=1.2, justify=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.first_line_indent = Inches(indent)
    pf.line_spacing = 1.0
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return rich(p, parts if isinstance(parts, list) else [(parts, "")])


def h1(num, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{num}.  {text.upper()}" if num else text.upper())
    r.font.name, r.font.size, r.bold = SERIF, Pt(9.5), False
    r.font.color.rgb = INK
    rPr = r._element.get_or_add_rPr()
    sc = OxmlElement("w:smallCaps")
    sc.set(qn("w:val"), "1")
    rPr.append(sc)
    return p


def h2(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name, r.font.size, r.italic = SERIF, Pt(9.5), True
    r.font.color.rgb = INK
    return p


def figure(path, caption_parts, width_in, full=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_together = True
    if os.path.exists(path):
        p.add_run().add_picture(str(path), width=Inches(width_in))
    else:
        r = p.add_run(f"[Figure: {os.path.basename(path)}]")
        r.italic = True
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if full else WD_ALIGN_PARAGRAPH.LEFT
    cp.paragraph_format.space_after = Pt(6)
    cp.paragraph_format.keep_together = True
    rich(cp, caption_parts, size=8.0)
    return p


def table_caption(parts, before=6):
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(before)
    cp.paragraph_format.space_after = Pt(2)
    cp.paragraph_format.keep_with_next = True
    rich(cp, parts, size=8.0)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), hexcolor)
    tcPr.append(sh)


def set_borders(tbl):
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        if edge in ("top", "bottom"):
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "12")
        elif edge == "insideH":
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "2")
        else:
            el.set(qn("w:val"), "none")
            el.set(qn("w:sz"), "0")
        el.set(qn("w:color"), "16181D" if edge in ("top", "bottom") else "C9CDD6")
        borders.append(el)
    tblPr.append(borders)


def make_table(headers, rows, widths=None, size=7.5, hl_row=None, align_left_col0=True):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_borders(t)

    # Set default table cell margins (padding)
    tblPr = t._tbl.tblPr
    tblCellMar = OxmlElement("w:tblCellMar")
    for edge, val in (("top", 40), ("bottom", 40), ("left", 60), ("right", 60)):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tblCellMar.append(node)
    tblPr.append(tblCellMar)

    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if (i == 0 and align_left_col0) else WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2.0)
        p.paragraph_format.space_before = Pt(2.0)
        r = p.add_run(h)
        r.font.name, r.font.size, r.bold = SANS, Pt(size), True
        r.font.color.rgb = INK
        shade(hdr[i], "F0F2F6")

    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if (i == 0 and align_left_col0) else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.space_before = Pt(1.5)
            mono = isinstance(v, str) and v.startswith(" ")
            txt = v[1:] if mono else v
            r = p.add_run(str(txt))
            r.font.name = "Consolas" if mono else SERIF
            r.font.size = Pt(size)
            r.font.color.rgb = INK
            if hl_row is not None and ri == hl_row:
                r.bold = True
        if hl_row is not None and ri == hl_row:
            for c in cells:
                shade(c, "E8ECF5")

    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                cell = row.cells[i]
                cell.width = Inches(w)
                tcPr = cell._tc.get_or_add_tcPr()
                tcW = tcPr.find(qn("w:tcW"))
                if tcW is None:
                    tcW = OxmlElement("w:tcW")
                    tcPr.append(tcW)
                tcW.set(qn("w:w"), str(int(w * 1440)))
                tcW.set(qn("w:type"), "dxa")

    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        if not trPr.findall(qn("w:cantSplit")):
            trPr.append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            for par in cell.paragraphs:
                par.paragraph_format.keep_together = True
    return t


# ══════════════════════════════════════════════════════════════════════════
# TITLE & AUTHOR BLOCK (Single-Column)
# ══════════════════════════════════════════════════════════════════════════
set_cols(s, 1)

para("SmartDetect: Face-Anchored Identity Arbitration for\nAuditable Cross-Camera Person Re-Identification",
     size=18, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)

AUTHORS_INFO = [
    ("Krishiv Kameshwar B. S.", ["Department of Information Technology", "Sri Krishna College of Technology", "Coimbatore, Tamil Nadu, India", "727823tuit113@skct.edu.in"]),
    ("Kishan S. G.", ["Department of Information Technology", "Sri Krishna College of Technology", "Coimbatore, Tamil Nadu, India", "727823tuit111@skct.edu.in"]),
    ("Lokeshkumar D.", ["Department of Information Technology", "Sri Krishna College of Technology", "Coimbatore, Tamil Nadu, India", "727823tuit115@skct.edu.in"]),
    ("K. Sinduja", ["Assistant Professor", "Department of Information Technology", "Sri Krishna College of Engineering and Technology", "Coimbatore, Tamil Nadu, India", "sindujak@skcet.ac.in"]),
]


def author_cell(cell, name, lines):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(name)
    r.font.name, r.font.size = SERIF, Pt(9.8)
    r.font.color.rgb = INK
    for i, ln in enumerate(lines):
        q = cell.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.space_after = Pt(0)
        q.paragraph_format.space_before = Pt(1 if i == 0 else 0)
        q.paragraph_format.line_spacing = 1.0
        rr = q.add_run(ln)
        rr.font.name, rr.font.size = SERIF, Pt(7.6)
        rr.italic = (i < len(lines) - 1)
        rr.font.color.rgb = MUT


def no_borders(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        b.append(el)
    tblPr.append(b)


t = doc.add_table(rows=1, cols=4)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
no_borders(t)
for i, info in enumerate(AUTHORS_INFO):
    author_cell(t.rows[0].cells[i], *info)
for cell in t.rows[0].cells:
    cell.width = Inches(1.81)

para(after=6)

# ══════════════════════════════════════════════════════════════════════════
# TWO-COLUMN BODY (Abstract, Introduction, etc.)
# ══════════════════════════════════════════════════════════════════════════
new_section(2)

CW = 3.40      # Column width
FW = 6.20      # Full width

ABSTRACT = (
    "Person re-identification (re-ID) is conventionally evaluated as ranked retrieval over closed, static galleries. "
    "In autonomous multi-camera surveillance, however, a system operates in an open-set, self-enrolling regime where unseen individuals "
    "are continuously observed, and every identification decision dynamically mutates the reference gallery for subsequent frames. "
    "Under this regime, unconstrained appearance cues (torso colour histograms and body descriptors) induce a critical failure mode: "
    "identity collapse, wherein distinct individuals are fused under a single identifier, contaminating photographic evidence logs. "
    "We present SmartDetect, a multi-camera tracking and re-ID framework governed by face-anchored identity arbitration. "
    "SmartDetect enforces an asymmetrical evidence hierarchy: quality-gated facial embeddings (ArcFace) alone create or veto identities, "
    "while clothing colour and body Re-ID (OSNet) are restricted to bounded temporal re-association. Three hardening mechanisms—a facial "
    "contradiction veto, an ID-switch contradiction guard, and a face-confirmed evidence gate—prevent gallery corruption. "
    "Evaluated on the ChokePoint portal benchmark (25 subjects, 3 cameras, 2908 person-frames), an unhardened baseline collapses into "
    "7 merged identities (0.0% purity, 15.5% precision at 99.8% coverage). SmartDetect substantially mitigates identity collapse under the evaluated ChokePoint protocol, "
    "achieving 99.7% identity-assignment precision, 93.2% purity, and 100.0% evidence precision at 98.6% coverage, adding only 6.8 ms per-frame CPU latency."
)

ap = doc.add_paragraph()
ap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
ap.paragraph_format.space_after = Pt(4)
ap.paragraph_format.line_spacing = 1.0
rich(ap, [("Abstract—", "bi"), (ABSTRACT, "b")], size=8.8)

kp = doc.add_paragraph()
kp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
kp.paragraph_format.space_after = Pt(4)
kp.paragraph_format.line_spacing = 1.0
rich(kp, [("Index Terms—", "bi"),
          ("Person re-identification, open-set recognition, face recognition, multi-object tracking, "
           "identity arbitration, biometric data governance, video surveillance.", "")], size=8.8)

# ── Section I: Introduction ──────────────────────────────────────────────
h1("I", "Introduction")
body("Person re-identification (re-ID) research has predominantly focused on ranked retrieval over static, curated galleries [1]. "
     "Standard metrics such as Rank-1 accuracy and mean average precision (mAP) evaluate a closed-set, stateless, read-only protocol: "
     "the query identity is assumed present in the gallery, probe evaluations do not alter gallery state, and outputs are rankings for human review.")

body([("Deployed surveillance systems violate all three assumptions simultaneously. First, the task is ", ""),
      ("open-set", "b"), (": un-enrolled pedestrians continuously enter camera views. Second, it is ", ""),
      ("stateful and self-enrolling", "b"), (": when the system fails to match an individual, it mints a new gallery profile (SDT-XXXX) that conditions all future matches. "
       "Third, it is ", ""), ("evidentiary", "b"), (": it persists timestamped sightings and photographic crops used in auditing and forensics.", "")])

body([("In this environment, relying on weak appearance cues (colour histograms or body-shape descriptors) causes a compounding failure: ", ""),
      ("identity collapse", "b"), (". When an individual's face is occluded or distant, appearance heuristics may falsely match another person in similar attire. "
       "Once associated, the stored feature template drifts toward a multi-person centroid, accelerating further false merges. "
       "Crucially, identity collapse is ", ""), ("invisible to Rank-n and mAP", "i"),
       (", which evaluate static retrieval rather than dynamic gallery contamination.", "")])

body("To address this, we present SmartDetect, an asynchronous multi-camera surveillance framework founded on face-anchored identity arbitration. "
     "The primary novelty of SmartDetect is not the introduction of a new neural backbone, but an asymmetrical evidence hierarchy: "
     "quality-gated facial embeddings act as the sole ground truth for identity creation and vetoes, while clothing colour and body Re-ID "
     "are restricted to short-term, time-bounded re-association.")

h2("A.  Key Contributions")
for lead, txt in [
    ("1) Face-Anchored Identity Arbitration:", " An asymmetrical evidence hierarchy where a quality-gated ArcFace embedding is necessary to mint an identity and sufficient to veto non-facial associations, treating an unmatched face as a decisive stranger verdict rather than a lookup failure."),
    ("2) Identity-Collapse Protection:", " Hardening mechanisms including a facial contradiction veto on appearance proposals and a 2-strike ID-switch guard countering tracker occlusion errors."),
    ("3) Evidence-Aware Persistence:", " A face-confirmed evidence gate decoupling transient on-screen annotations from persistent database writes, preventing false photographic evidence logging."),
    ("4) Open-Set Evaluation Methodology:", " An exhaustive outcome partition (CORRECT, CONTAMINATED, UNASSIGNED), separating precision from coverage without artificial F-score blending, paired with a dataset adequacy gate."),
    ("5) Ablation-Based Validation:", " Rigorous quantification on the ChokePoint benchmark demonstrating that SmartDetect substantially mitigates identity collapse under the evaluated ChokePoint protocol with 99.7% identity-assignment precision and 100.0% evidence precision.")
]:
    body([(lead, "b"), (txt, "")])

# ── Section II: Related Work ─────────────────────────────────────────────
h1("II", "Related Work")
h2("A.  Closed-Set Retrieval vs. Open-Set Re-ID")
body("Deep metric learning, part-based pooling, and attention networks [1]–[4] have pushed closed-set Rank-1 retrieval accuracy above 90% on benchmarks such as Market-1501 and DukeMTMC. "
     "However, these models assume pre-cropped boxes and static galleries. Open-set re-ID formulations [5] acknowledge novel queries but evaluate single-probe rejection rather than the cumulative, "
     "self-corrupting dynamics of autonomous multi-camera enrollment.")

h2("B.  Appearance Re-ID vs. Face Recognition")
body("Face recognition backbones like ArcFace [8] provide highly discriminative 512-d embeddings under sufficient resolution but fail when subjects turn away. "
     "Conversely, body Re-ID backbones like OSNet [6] capture global morphology but struggle with inter-class visual similarities (e.g., uniform clothing). "
     "SmartDetect explicitly structures their interaction, preventing appearance features from overriding or minting facial identities.")

h2("C.  Multi-Camera Tracking and Spatio-Temporal Priors")
body("Multi-target multi-camera tracking (MTMCT) often uses spatial-temporal topology to constrain inter-camera travel times [1], [11]. "
     "While valuable as filters, topological priors cannot assert identity independently without risking false links across pedestrians traversing identical paths.")

h2("D.  Novelty Gap")
body("Existing literature focuses almost exclusively on feature representation. The unaddressed gap is the cumulative failure of autonomous self-enrollment and gallery corruption. "
     "SmartDetect bridges this gap through structural evidence gating.")

# ── Section III: SmartDetect System Architecture ────────────────────────
h1("III", "SmartDetect System Architecture")

figure(FIGS / "fig_pipeline.png",
       [("Fig. 1.  ", "b"),
        ("SmartDetect per-frame decision pipeline. Stages 5–8 enforce face-anchored arbitration and evidence gating.", "")], CW)

h2("A.  Decoupled Threading Ingest Engine")
body("To maintain high-throughput video delivery on commodity CPU hardware, SmartDetect decouples video capture from deep inference into an asynchronous architecture. "
     "A per-camera capture thread ingests frames at native rates (up to 30 FPS), pushes the newest frame into a depth-1 drop-oldest queue, "
     "and streams annotated MJPEG video using cached tracklet states. An asynchronous ML worker thread executes the full analysis pipeline (62.4 ms/frame, or ~16 FPS continuous single-camera throughput, paced at 1–3 Hz for multi-stream CPU workloads), "
     "ensuring stream smoothness is decoupled from heavy inference latency.")

h2("B.  Modular Vision Backbones")
body("SmartDetect integrates proven, pretrained modular backbones without fine-tuning: "
     "(1) YOLOv8n person detection with resolution-aware input sizing (416px for webcam, 640px for HD, 960px for 4K); "
     "(2) ByteTrack [9] for multi-object tracking; (3) InsightFace ArcFace (`buffalo_l`) [8] for 512-d facial embeddings via ONNX Runtime; "
     "(4) OSNet x1.0 [6] for 512-d body Re-ID descriptors; and (5) K-means HSV torso colour clustering.")

h2("C.  Implementation Configuration")
body("Table I summarizes the operational parameters and thresholds configured across the system.", indent=0)

table_caption([("TABLE I.    Implementation and Arbitration Configuration", "b")])
make_table(["Subsystem / Parameter", "Value / Threshold", "Operational Role"],
           [["Detection (YOLOv8n)", "416 / 640 / 960 px", "Resolution-scaled person detection"],
            ["Tracking (ByteTrack)", "High-score association", "Temporal tracklet continuity"],
            ["Face Quality Gate", "Height ≥ 48px, Det ≥ 0.60", "Admission to biometric arbitration"],
            ["Face Match Threshold", "cos(sim) ≥ 0.56", "Primary ArcFace identity match"],
            ["Face Veto Threshold", "cos(sim) < 0.45", "Rejects contradictory appearance links"],
            ["ID-Switch Guard", "2 strikes (cos < 0.35)", "Invalidates cached track identities"],
            ["Evidence Gate", "cos(sim) ≥ 0.45", "Face-confirmed sighting persistence"],
            ["Torso Colour Window", "HSV dist ≤ 30.0 (10 min)", "Short-term temporal re-association"],
            ["Body Re-ID Window", "cos(sim) ≥ 0.68 (12 h)", "Same-day clothing re-association"],
            ["Multi-View Gallery", "Cap 5 views ([0.60, 0.85])", "Adaptive pose template memory"]],
           widths=[1.25, 1.15, 1.00])

# ── Section IV: Identity Arbitration and Evidence Hierarchy ──────────────
h1("IV", "Identity Arbitration and Evidence Hierarchy")

h2("A.  The Evidence Hierarchy")
body("Arbitration evaluates facial embedding f, torso HSV colour c, and OSNet descriptor r under four formal invariants:", indent=0)

for lead, txt in [
    ("H1 (Enrolment Anchor):", " Only a quality-gated face embedding f can mint a new SDT-XXXX identity. Secondary cues (c, r) cannot enroll individuals."),
    ("H2 (Stranger Decisiveness):", " A gate-passing face f matching no gallery entry is definitively classified as a stranger. Cues c and r are completely suppressed."),
    ("H3 (Bounded Re-association):", " Secondary cues (c, r) may only re-associate previously enrolled profiles within strict temporal windows (10 min for colour, 12 h for Re-ID), subject to immediate facial veto if f is present."),
    ("H4 (Evidence Confirmation):", " Persistent database sightings and photo crops require same-cycle facial confirmation matching the assigned code.")
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.space_after = Pt(1.0)
    rich(p, [(lead, "ba"), (" " + txt, "")])

h2("B.  Quality Gate and Face Veto")
body("Faces shorter than 48 px or with detector confidence <0.60 are excluded from arbitration, as sub-48px ArcFace embeddings exhibit severe interpolation distortion. "
     "When a secondary cue proposes an enrolled candidate p, if the subject possesses a visible face f, the face veto computes cosine similarity against p's stored template. "
     "If similarity falls below 0.45, the proposal is rejected outright.")

h2("C.  ByteTrack ID-Switch Guard")
body("During spatial intersections and occlusions, ByteTrack can swap track IDs. Because arbitration is cached per track for CPU efficiency, "
     "a tracker ID swap would permanently hijack an identity. The ID-switch guard (Algorithm 1) evaluates gate-passing faces against the cached profile, "
     "dropping the cache after 2 consecutive contradictions (<0.35 similarity).")

ALGO = [
    "Input: Track t, Cached identity ID_t, Current face embedding f",
    "State: Contradiction strikes S[t] (initially 0)",
    "if f is NULL or ID_t is NULL then return ID_t",
    "s = cosine_similarity(f, Gallery[ID_t])",
    "if s < 0.35 then",
    "    S[t] = S[t] + 1",
    "    if S[t] >= 2 then drop_cache(t); return NULL",
    "else",
    "    S[t] = 0   // Reset on agreement",
    "return ID_t",
]
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(4)
p.paragraph_format.space_after = Pt(1)
r = p.add_run("Algorithm 1: ByteTrack ID-Switch Contradiction Guard")
r.font.name, r.font.size, r.bold = SANS, Pt(8.0), True
for ln in ALGO:
    lp = doc.add_paragraph()
    lp.paragraph_format.space_after = Pt(0)
    lp.paragraph_format.line_spacing = 1.0
    lp.paragraph_format.left_indent = Inches(0.08)
    lr = lp.add_run(ln if ln else " ")
    lr.font.name, lr.font.size = "Consolas", Pt(7.0)
    lr.font.color.rgb = INK
para(after=4)

h2("D.  Evidence Gating and Registration Pose Gate")
body("SmartDetect strictly separates display labels from durable persistence. A track lacking an active face may retain its display annotation for visual continuity, "
     "but no database sighting or photographic crop is written to disk without same-cycle face confirmation (sim ≥ 0.45). "
     "Furthermore, the registration pose gate evaluates landmark ratios to reject extreme head poses (|yaw| > 0.35, pitch outside [0.35, 0.70]):", indent=0)

for eq, num in [("yaw = (n_x - e_x) / d", "(1)"),
                ("pitch = (n_y - e_y) / (m_y - e_y)", "(2)")]:
    ep = doc.add_paragraph()
    ep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ep.paragraph_format.space_before = Pt(2)
    ep.paragraph_format.space_after = Pt(2)
    r = ep.add_run(eq); r.font.name, r.font.size, r.italic = SERIF, Pt(9.0), True
    r2 = ep.add_run("    " + num); r2.font.name, r2.font.size = SERIF, Pt(8.5)

where_p = doc.add_paragraph()
where_p.paragraph_format.space_after = Pt(2)
rich(where_p, [("where ", "i"), ("e", "m"), (" is the inter-ocular midpoint, ", ""), ("d", "m"),
               (" is inter-ocular distance, ", ""), ("n", "m"), (" is the nose tip, and ", ""),
               ("m", "m"), (" is the mouth midpoint.", "")], size=8.5)

h2("E.  Split-over-Merge Asymmetry")
body("Faced with ambiguity, SmartDetect deliberately splits identities rather than merging them. "
     "A duplicate identity (two codes for one person) is localized, visible, and safely resolvable via operator merge tooling. "
     "A false merge (two people under one code) is silent, permanently corrupts the gallery template, and is irreversible.")

# ── Section V: Experimental Setup ────────────────────────────────────────
h1("V", "Experimental Setup")

h2("A.  Benchmark Dataset and Ground Truth Protocol")
body("We evaluate on the public ChokePoint portal dataset [7] (Sequence P1E_S1, 3 cameras, 25 subjects, 2908 ground-truth person-frames across 6876 video frames). "
     "Ground truth provides eye coordinates and subject IDs. The unit of measurement is the ground-truth person-frame (including detector misses). "
     "Because the contribution targets stateful identity-gallery contamination rather than benchmark retrieval accuracy, the evaluation focuses on a controlled multi-camera sequence where frame-level identity assignments and cross-camera observations can be explicitly audited. "
     "The evaluation is limited to the P1E_S1 sequence of the ChokePoint dataset and therefore does not establish generalization across portals, sequences, or datasets.")

h2("B.  Outcome Partition and Evaluation Metrics")
body("Let maj(c) denote the majority ground-truth identity assigned to code c. Every frame record r maps into an exhaustive, mutually exclusive partition:", indent=0)

for lead, txt in [
    ("• UNASSIGNED:", " No code issued ('Detecting...' sentinel or detector miss)."),
    ("• CORRECT:", " maj(c) equals the ground-truth subject."),
    ("• CONTAMINATED:", " maj(c) does not match ground truth (identity collision).")
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.space_after = Pt(1.0)
    rich(p, [(lead, "b"), (" " + txt, "")])

body("Precision, Coverage, and Evidence Precision are defined as:", indent=0)
for eq, num in [("Precision = n_correct / (n_correct + n_contaminated)", "(3)"),
                ("Coverage = (n_correct + n_contaminated) / n_total", "(4)"),
                ("Evidence Precision = n_correct_persisted / n_total_persisted", "(5)")]:
    ep = doc.add_paragraph()
    ep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ep.paragraph_format.space_before = Pt(2)
    ep.paragraph_format.space_after = Pt(2)
    r = ep.add_run(eq); r.font.name, r.font.size, r.italic = SERIF, Pt(9.0), True
    r2 = ep.add_run("    " + num); r2.font.name, r2.font.size = SERIF, Pt(8.5)

body("Proportions are reported with 95% Wilson score confidence intervals [10]. Additional metrics include Identity Purity (fraction of codes mapping to exactly 1 person), "
     "Duplicates per Person, Tracker ID Switches, and Cross-Camera Re-association Rate.")

h2("C.  Ablation Configurations")
body("We evaluate four cumulative configurations: "
     "(A) Pre-hardening Baseline: unconstrained colour/Re-ID fallback without face anchoring; "
     "(B) Face Anchor: adds H2 face anchoring and face veto; "
     "(C) + ID-Switch Guard: adds 2-strike contradiction cache invalidation; and "
     "(D) Full Pipeline: adds face-confirmed evidence gating.")

# ── Section VI: Results and Ablation Study ───────────────────────────────
h1("VI", "Results and Ablation Study")

# ── Full-Width Section for Table II and Fig. 2 ───────────────────────────
new_section(1)
table_caption([("TABLE II.    HEADLINE EVALUATION METRICS ON CHOKEPOINT P1E_S1 (2908 SCORED FRAMES)", "b")])
make_table(["Configuration", "Precision", "Coverage", "Purity", "Contam. Frames", "Evidence Precision", "IDs", "Cross-Camera"],
           [["A – Pre-hardening", "15.5%", "99.8%", "0.0%", "2454", "—*", "7", "8.0%"],
            ["B – Face Anchor", "53.8%", "98.6%", "48.3%", "1324", "—*", "29", "28.0%"],
            ["C – + ID-Switch Guard", "99.7%", "98.6%", "93.2%", "9", "—*", "59", "50.0%"],
            ["D – Full Pipeline", "99.7%", "98.6%", "93.2%", "9", "100.0%", "59", "50.0%"]],
           widths=[1.45, 0.70, 0.70, 0.65, 0.95, 1.05, 0.60, 1.00],
           size=7.8, hl_row=3)
para("*Evidence Precision is not applicable (—) for Configs A–C because face-confirmed evidence gating is disabled; in Config D, gating prevents transient contaminated frames from writing to disk (0 incorrect rows persisted).", size=7.2, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)

figure(FIGS / "fig_collapse.png",
       [("Fig. 2.  ", "b"),
        ("Identity collapse visualization on ChokePoint P1E_S1. (a) In Config A, 7 identifiers smear across 25 people with 0.0% purity. "
         "(b) In Config D, SmartDetect enforces clean diagonal separation with 93.2% purity; residual off-diagonal mass represents safe over-splitting.", "")],
       FW, full=True)
new_section(2)

h2("A.  Ablation Analysis")
body([("1) Reproduction of Identity Collapse (Config A): ", "b"),
      ("The unhardened baseline exhibits catastrophic collapse, minting only 7 identifiers for 25 individuals (0.0% purity). "
       "The dominant identifier absorbed 22 distinct people across 2454 contaminated frames. While achieving 99.8% coverage, identity-assignment precision was only 15.5% (Table II, Fig. 2a).", "")])

body([("2) Impact of Face-Anchoring (Config B): ", "b"),
      ("Enforcing face-anchored matching triples precision to 53.8% and increases purity to 48.3% (29 IDs minted), reducing contaminated frames from 2454 to 1324. "
       "However, tracker occlusion swaps continue to cause lingering contamination on cached tracks.", "")])

body([("3) Decisive Role of ID-Switch Guard (Config B → Config C): ", "b"),
      ("Adding contradiction invalidation produces the single most dramatic ablation improvement: contaminated frames plummet from 1324 to just 9 (a 99.3% reduction in contamination). "
       "This raises identity-assignment precision from 53.8% to 99.7% and purity from 48.3% to 93.2%, while cross-camera re-association increases from 28.0% to 50.0%. "
       "As predicted, raw tracker ID switches rose from 44 to 57, confirming that correcting hijacked tracks registers as switches.", "")])

body([("4) Zero-Contamination Evidence Persistence (Config D): ", "b"),
      ("Config D's evidence gate achieves 100.0% evidence precision, completely filtering out the 9 transient contradiction frames so that 0 incorrect sighting rows are written to disk (2402 clean sighting rows persisted). "
       "Total pipeline latency increases by only 6.8 ms/frame (55.6 ms to 62.4 ms), representing a modest 12% computational overhead.", "")])

figure(FIGS / "fig_buckets.png",
       [("Fig. 3.  ", "b"),
        ("Ground-truth outcome partition across Configurations A–D (n = 2908 frames).", "")], CW)

# ── Section VII: Data Governance and Technical Privacy ───────────────────
h1("VII", "Data Governance and Technical Privacy")
body("Surveillance biometrics demand active technical enforcement. SmartDetect implements technical safeguards aligned with privacy and data-governance requirements [13], [14]:", indent=0)

for lead, txt in [
    ("• Consent-Driven Retention TTLs:", " Enrolled profiles are classified as 'consented' (365d TTL), 'dataset' (3650d TTL), or 'unknown' (7d automated purge TTL)."),
    ("• Verified Right-to-be-Forgotten:", " Deletion requests purge database records, physical photo crops, and active in-memory camera thread caches, verifying disk/DB state before issuing an immutable ErasureReceipt."),
    ("• Append-Only Biometric Audit Logging:", " All biometric searches and lookups are logged in an immutable ledger (AuditLog) recording actor, timestamp, query image, and outcome.")
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.space_after = Pt(1.0)
    rich(p, [(lead, "b"), (" " + txt, "")])

# ── Section VIII: Discussion and Limitations ─────────────────────────────
h1("VIII", "Discussion and Limitations")
body([("1) Sequence and Dataset Scope: ", "b"),
      ("The evaluation is limited to the P1E_S1 sequence of the ChokePoint dataset and therefore does not establish generalization across portals, sequences, or datasets.", "")])
body([("2) Proof-of-Concept Protocol: ", "b"),
      ("This sequence is used as a controlled proof-of-concept protocol for auditing stateful identity assignments rather than as evidence of dataset-wide generalization.", "")])
body([("3) Face Visibility Requirement: ", "b"),
      ("Face visibility is strictly required for reliable identity arbitration; subjects without visible faces remain unassigned or transiently tracked, trading wide-area crowd coverage for zero contaminated persistent evidence under the evaluated protocol.", "")])
body([("4) Face Resolution Constraints: ", "b"),
      ("Faces with height below 48px or detector confidence < 0.60 are rejected from biometric arbitration to prevent poor-quality embeddings.", "")])
body([("5) Duplicate Identity Trade-off: ", "b"),
      ("Under the intentional split-over-merge asymmetry, duplicate identity generation is accepted as a safety trade-off to prevent irreversible gallery contamination.", "")])
body([("6) Duplicate Identity Metric: ", "b"),
      ("The evaluation yielded an average of 1.36 duplicate codes per person, manageable via operator duplicate-suggestion merge tooling.", "")])
body([("7) Cross-Camera Re-Association Ceiling: ", "b"),
      ("Without spatial camera graph topology, cross-camera re-association is bounded to 50.0%.", "")])
body([("8) Single-Node CPU Latency: ", "b"),
      ("Per-frame analysis latency (62.4 ms/frame, ~16 FPS continuous single-camera throughput) limits concurrent high-resolution streams on a single CPU node.", "")])
body([("9) Multi-Stream CPU Pacing: ", "b"),
      ("Multi-stream CPU workloads are paced at 1–3 Hz to prevent queue starvation and maintain stream smoothness.", "")])
body([("10) Baseline Scope: ", "b"),
      ("A direct empirical comparison against conventional appearance-only multi-camera Re-ID baselines remains an important direction for future evaluation.", "")])

# ── Section IX: Conclusion and Future Work ───────────────────────────────
h1("IX", "Conclusion and Future Work")
body("Ranked retrieval benchmarks fail to expose identity collapse in autonomous surveillance. "
     "SmartDetect demonstrates that face-anchored identity arbitration, contradiction guards, and evidence gating substantially mitigate identity collapse under the evaluated ChokePoint protocol "
     "(99.7% identity-assignment precision, 93.2% purity, 100.0% evidence precision) at minimal CPU cost (62.4 ms/frame). "
     "Future work includes evaluating direct appearance-only baselines across broader portal sequences, tracklet-level consensus voting, and spatial-topological graph constraints.")

# ── Section X: Acknowledgment ────────────────────────────────────────────
h1("", "Acknowledgment")
body("The authors thank the Department of Information Technology at Sri Krishna College of Technology and Sri Krishna College of Engineering and Technology for supporting this research, "
     "and the creators of the ChokePoint dataset for providing annotated multi-camera surveillance benchmarks.", indent=0)

# ── Section XI: References ───────────────────────────────────────────────
h1("", "References")
REFS = [
    "H. Wang, H. Du, Y. Zhao, and J. Yan, \"A comprehensive overview of person re-identification approaches,\" IEEE Access, vol. 8, pp. 45556–45583, 2020.",
    "Y. Sun, L. Zheng, Y. Yang, Q. Tian, and S. Wang, \"Beyond part models: Person retrieval with refined part pooling,\" in Proc. ECCV, 2018, pp. 501–518.",
    "W. Li, X. Zhu, and S. Gong, \"Harmonious attention network for person re-identification,\" in Proc. IEEE/CVF CVPR, 2018, pp. 2285–2294.",
    "A. Hermans, L. Beyer, and B. Leibe, \"In defense of the triplet loss for person re-identification,\" arXiv:1703.07737, 2017.",
    "Z. Zheng, L. Zheng, Z. Hu, and Y. Yang, \"Open set adversarial examples,\" arXiv:1809.02681, 2018.",
    "K. Zhou, Y. Yang, A. Cavallaro, and T. Xiang, \"Omni-scale feature learning for person re-identification,\" in Proc. IEEE/CVF ICCV, 2019, pp. 3702–3712.",
    "Y. Wong, S. Chen, S. Mau, C. Sanderson, and B. C. Lovell, \"Patch-based Probabilistic Image Quality Assessment for Face Selection and Improved Video-based Face Recognition,\" in Proc. IEEE Computer Society Conference on Computer Vision and Pattern Recognition Workshops (CVPRW), Colorado Springs, CO, USA, 2011, pp. 74–81, doi: 10.1109/CVPRW.2011.5981881.",
    "J. Deng, J. Guo, N. Xue, and S. Zafeiriou, \"ArcFace: Additive angular margin loss for deep face recognition,\" in Proc. IEEE/CVF CVPR, 2019, pp. 4690–4699.",
    "Y. Zhang et al., \"ByteTrack: Multi-object tracking by associating every detection box,\" in Proc. ECCV, 2022, pp. 1–21.",
    "E. B. Wilson, \"Probable inference, the law of succession, and statistical inference,\" J. Amer. Statist. Assoc., vol. 22, no. 158, pp. 209–212, 1927.",
    "E. Ristani et al., \"Performance measures and a data set for multi-target, multi-camera tracking,\" in Proc. ECCV Workshops, 2016, pp. 17–35.",
    "L. Zheng et al., \"Scalable person re-identification: A benchmark,\" in Proc. IEEE ICCV, 2015, pp. 1116–1124.",
    "Government of India, The Digital Personal Data Protection Act, 2023, Act No. 22 of 2023.",
    "European Parliament and Council, Regulation (EU) 2016/679 (GDPR), Art. 9 & 17, 2016."
]

for i, ref in enumerate(REFS, 1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(1.2)
    p.paragraph_format.left_indent = Inches(0.20)
    p.paragraph_format.first_line_indent = Inches(-0.20)
    r = p.add_run(f"[{i}]  ")
    r.font.name, r.font.size = SERIF, Pt(7.8)
    r2 = p.add_run(ref)
    r2.font.name, r2.font.size = SERIF, Pt(7.8)
    r2.font.color.rgb = INK

doc.save(OUT)
print(f"Successfully generated: {OUT}")
print(f"File size: {round(OUT.stat().st_size / 1024, 1)} KB")
